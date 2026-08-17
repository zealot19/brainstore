import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


def render_xlsx(result) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapport"
    ws.append([result.title])
    ws["A1"].font = Font(bold=True, size=14)
    if result.period_label:
        ws.append([result.period_label])
    ws.append([c.label for c in result.columns])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for g in result.groups:
        if g.label:
            ws.append([g.label])
            ws[ws.max_row][0].font = Font(bold=True, italic=True)
        for row in g.rows:
            ws.append([row.get(c.key) for c in result.columns])
        if g.subtotal:
            ws.append([g.subtotal.get(c.key) for c in result.columns])
    if result.totals:
        ws.append([result.totals.get(c.key) for c in result.columns])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, c in enumerate(result.columns, 1):
        if c.align == "right":
            for col_cells in ws.iter_cols(min_col=i, max_col=i):
                for cc in col_cells:
                    cc.alignment = Alignment(horizontal="right")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
